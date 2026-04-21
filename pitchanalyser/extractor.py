"""
extractor.py - Extract frames from videos and pages from pitch decks
All processing is done locally. No data leaves the machine.
"""

import os
import subprocess
import base64
from pathlib import Path


def extract_frames_from_video(
    video_path: str, output_dir: str, frames_per_minute: int = 2
) -> list[str]:
    """
    Extract frames from a video file using ffmpeg.
    Returns list of base64-encoded image strings.

    Args:
        video_path: Path to the video file
        output_dir: Directory to save extracted frames
        frames_per_minute: How many frames to extract per minute of video

    Returns:
        List of base64-encoded JPEG image strings
    """
    os.makedirs(output_dir, exist_ok=True)

    # Calculate frame rate (frames per second)
    fps = frames_per_minute / 60.0

    output_pattern = os.path.join(output_dir, "frame_%04d.jpg")

    cmd = [
        "ffmpeg",
        "-i",
        video_path,
        "-vf",
        f"fps={fps},scale=1280:-1",  # Scale width to 1280px, maintain aspect ratio
        "-q:v",
        "3",  # JPEG quality (2=best, 31=worst)
        "-y",  # Overwrite existing files
        output_pattern,
    ]

    try:
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=300,  # 5 minute timeout
        )
        if result.returncode != 0:
            print(f"    WARNING: ffmpeg error: {result.stderr[-500:]}")
    except FileNotFoundError:
        raise RuntimeError(
            "ffmpeg not found. Install it with:\n"
            "  macOS:  brew install ffmpeg\n"
            "  Ubuntu: sudo apt install ffmpeg\n"
            "  Windows: https://ffmpeg.org/download.html"
        )
    except subprocess.TimeoutExpired:
        print("    WARNING: ffmpeg timed out after 5 minutes")

    # Load and encode frames
    frame_files = sorted(Path(output_dir).glob("frame_*.jpg"))
    encoded_frames = []
    for frame_file in frame_files:
        with open(frame_file, "rb") as f:
            encoded = base64.b64encode(f.read()).decode("utf-8")
            encoded_frames.append(encoded)

    return encoded_frames


def extract_pages_from_deck(
    deck_path: str, output_dir: str, dpi: int = 150
) -> list[str]:
    """
    Extract pages from a PDF or PPTX pitch deck as images.
    Returns list of base64-encoded image strings.

    Args:
        deck_path: Path to the deck file (.pdf or .pptx)
        output_dir: Directory to save extracted page images
        dpi: Resolution for PDF rendering (higher = better quality, slower)

    Returns:
        List of base64-encoded PNG image strings
    """
    os.makedirs(output_dir, exist_ok=True)
    ext = Path(deck_path).suffix.lower()

    if ext == ".pdf":
        return _extract_pdf_pages(deck_path, output_dir, dpi)
    elif ext in (".pptx", ".ppt"):
        return _extract_pptx_pages(deck_path, output_dir, dpi)
    else:
        raise ValueError(f"Unsupported deck format: {ext}")


def _extract_pdf_pages(pdf_path: str, output_dir: str, dpi: int) -> list[str]:
    """Extract PDF pages as images using pymupdf."""
    try:
        import fitz  # pymupdf
    except ImportError:
        raise RuntimeError("pymupdf not found. Install it with:\n  pip install pymupdf")

    doc = fitz.open(pdf_path)
    encoded_pages = []
    zoom = dpi / 72.0  # PDF default is 72 DPI
    mat = fitz.Matrix(zoom, zoom)

    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=mat)
        img_path = os.path.join(output_dir, f"page_{i + 1:03d}.png")
        pix.save(img_path)

        with open(img_path, "rb") as f:
            encoded = base64.b64encode(f.read()).decode("utf-8")
            encoded_pages.append(encoded)

    doc.close()
    return encoded_pages


def _extract_pptx_pages(pptx_path: str, output_dir: str, dpi: int) -> list[str]:
    """
    Extract PPTX slides as images.
    Strategy: Convert to PDF first via LibreOffice, then extract pages.
    Falls back to text extraction if LibreOffice is not available.
    """
    # Try LibreOffice conversion first
    pdf_path = os.path.join(output_dir, "converted.pdf")
    try:
        result = subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                output_dir,
                pptx_path,
            ],
            capture_output=True,
            text=True,
            timeout=120,
        )
        # LibreOffice outputs with same stem as input
        stem = Path(pptx_path).stem
        converted = os.path.join(output_dir, f"{stem}.pdf")
        if os.path.exists(converted):
            return _extract_pdf_pages(converted, output_dir, dpi)
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    # Fallback: extract slide text as a simple text representation
    print(
        "    WARNING: LibreOffice not found. Extracting text from PPTX instead of images."
    )
    print("    For better results, install LibreOffice: https://www.libreoffice.org/")
    return _extract_pptx_as_text_images(pptx_path, output_dir)


def _extract_pptx_as_text_images(pptx_path: str, output_dir: str) -> list[str]:
    """Fallback: render PPTX slide text onto simple images using Pillow."""
    try:
        from pptx import Presentation
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        raise RuntimeError(
            "python-pptx and/or Pillow not found. Install with:\n"
            "  pip install python-pptx Pillow"
        )

    prs = Presentation(pptx_path)
    encoded_pages = []

    for i, slide in enumerate(prs.slides):
        # Gather all text from the slide
        lines = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lines.append(shape.text.strip())

        slide_text = "\n".join(lines) if lines else "(No text on slide)"

        # Render text onto an image
        img = Image.new("RGB", (1280, 720), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)

        # Try to use a basic font, fall back to default
        try:
            font = ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 24
            )
            small_font = ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 18
            )
        except:
            font = ImageFont.load_default()
            small_font = font

        draw.text((40, 20), f"Slide {i + 1}", fill=(100, 100, 100), font=small_font)
        draw.text((40, 60), slide_text[:2000], fill=(0, 0, 0), font=font)

        img_path = os.path.join(output_dir, f"slide_{i + 1:03d}.png")
        img.save(img_path)

        with open(img_path, "rb") as f:
            encoded = base64.b64encode(f.read()).decode("utf-8")
            encoded_pages.append(encoded)

    return encoded_pages
